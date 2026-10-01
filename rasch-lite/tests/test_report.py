"""Ekspor: workbook Excel, laporan HTML mandiri, laporan PDF, folder grafik."""

import os
import re

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from raschlite import theme
from raschlite.core.analysis import run_analysis
from raschlite.core.data import prepare_data, read_table
from raschlite.gui import strings as S
from raschlite.interpret import interpret, rules
from raschlite.report import export_excel, export_html
from raschlite.report.html import chart_plan
from raschlite.resources import sample_path


def _analyse(kind, model):
    df = read_table(sample_path(kind))
    items = [c for c in df.columns if c not in ("ID", "Jenis_Kelamin")]
    res = run_analysis(prepare_data(df, items, id_col="ID", group_col="Jenis_Kelamin"), model)
    return res, interpret(res)


@pytest.fixture(scope="module")
def dicho():
    return _analyse("dichotomous", "dichotomous")


@pytest.fixture(scope="module")
def pcm():
    return _analyse("polytomous", "pcm")


def test_draba_t_threshold_depends_on_number_of_items(dicho):
    assert rules.dif_t_min(20) == 2.0 and rules.dif_t_min(21) == 2.4
    res, _ = dicho
    assert res.settings["dif_t_min"] == 2.0
    flagged = res.dif[res.dif["flag_dif"]]
    assert (flagged["contrast"].abs() >= rules.DIF_CONTRAST_MIN).all() and (flagged["t"].abs() > 2.0).all()


def test_excel_has_one_sheet_per_table_with_numeric_cells(pcm, tmp_path):
    res, interp = pcm
    path = export_excel(res, interp, tmp_path / "hasil.xlsx", "contoh_politomus.csv")
    wb = load_workbook(path)
    for key in ("summary", "items", "persons", "categories", "dif", "loadings", "q3", "q3_pairs", "notes",
                "iterations", "settings"):
        assert S.SHEETS[key] in wb.sheetnames, key
    ws = wb[S.SHEETS["items"]]
    headers = [c.value for c in ws[1]]
    assert headers[0] == S.COLUMNS["item"] and "Infit MNSQ" in headers and "Measure (logit)" in headers
    assert ws.max_row == len(res.items) + 1
    measure_col = headers.index("Measure (logit)") + 1
    assert isinstance(ws.cell(row=2, column=measure_col).value, float)
    assert ws.cell(row=2, column=measure_col).value == pytest.approx(res.items["measure"].iloc[0])
    cat = wb[S.SHEETS["categories"]]
    flagged_fill = {cat.cell(row=r, column=1).fill.fgColor.rgb for r in range(2, cat.max_row + 1)}
    assert any(str(c).endswith(theme.FLAG_ROW.lstrip("#").upper()) for c in flagged_fill)
    summary_text = " ".join(str(c.value) for row in wb[S.SHEETS["summary"]].iter_rows() for c in row if c.value)
    for light in interp.lights:
        assert light.sentence in summary_text


def test_html_is_self_contained_and_complete(pcm, tmp_path):
    res, interp = pcm
    path = export_html(res, interp, tmp_path / "laporan.html", "contoh_politomus.csv")
    text = path.read_text(encoding="utf-8")
    assert not re.search(r'(src|href)="(https?:)?//', text), "ada sumber daya eksternal"
    assert "<script" not in text
    plan, all_items = chart_plan(res, 30)
    assert all_items and text.count("data:image/svg+xml;base64,") == len(plan)
    for light in interp.lights:
        assert light.title in text
    for heading in (S.SUMMARY_1MIN, S.EXPLANATION, S.REPORT_CHARTS, S.REPORT_TABLES, S.REPORT_GLOSSARY):
        assert heading in text
    assert text.count(S.HOW_TO_READ) == len(plan)


def test_html_escapes_user_supplied_names(tmp_path):
    rng = np.random.default_rng(3)
    X = (rng.random((200, 6)) < 0.5).astype(int)
    df = pd.DataFrame(X, columns=["<b>Q1</b>", "Q2", "Q3", "Q4", "Q5", "Q6"])
    res = run_analysis(prepare_data(df, list(df.columns)), "dichotomous")
    text = export_html(res, interpret(res), tmp_path / "x.html").read_text(encoding="utf-8")
    assert "<b>Q1</b>" not in text and "&lt;b&gt;Q1&lt;/b&gt;" in text


def test_pdf_and_full_export_folder(dicho, tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    pytest.importorskip("PySide6.QtGui", exc_type=ImportError)
    from raschlite.report import export_all

    res, interp = dicho
    steps = []
    written = export_all(res, interp, tmp_path, "contoh_dikotomus.csv", progress=lambda k, n, m: steps.append((k, n)))
    assert steps[-1][0] == steps[-1][1]
    for key in ("excel", "html", "pdf"):
        assert written[key].exists() and written[key].stat().st_size > 0
    pdf = written["pdf"].read_bytes()
    assert pdf.startswith(b"%PDF")
    assert len(re.findall(rb"/Type\s*/Page[^s]", pdf)) >= 10
    n_items = len(res.coded.item_names)
    n_global = 6 + int(res.dif is not None)  # grafik non-per-butir data dikotomus (+ DIF)
    assert len(written["charts"]) == 2 * (n_global + n_items)
    assert {p.suffix for p in written["charts"]} == {".png", ".svg"}
