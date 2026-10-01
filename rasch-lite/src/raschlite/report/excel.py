"""Ekspor workbook Excel: satu sheet per tabel (openpyxl, tanpa Qt).

Angka disimpan sebagai angka (bukan teks) agar dapat diolah ulang; tampilan
desimal diatur lewat format sel. Baris yang perlu diperiksa diberi warna latar.
"""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import theme
from ..gui import strings as S
from ..interpret.narrative import STATUS_LABELS
from .common import DECIMALS, column_label, dif_labels, generated_line, metadata, table_frame

def _hex(color: str) -> str:
    return color.lstrip("#").upper()


HEADER_FILL = PatternFill("solid", fgColor=_hex(theme.SAND))
FLAG_FILL = PatternFill("solid", fgColor=_hex(theme.FLAG_ROW))
MUTED_FILL = PatternFill("solid", fgColor=_hex(theme.BACKGROUND))
BOLD = Font(bold=True, color=_hex(theme.INK))
STATUS_FILLS = {k: _hex(v) for k, v in theme.STATUS_TINT.items()}


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return S.YES if v else ""
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return float(v) if math.isfinite(float(v)) else None
    return v if isinstance(v, (int, str)) else str(v)


def _number_format(decimals: int) -> str:
    return "0" if decimals == 0 else "0." + "0" * decimals


def _write_rows(ws, headers: list[str], rows, start_row: int = 1, decimals: list[int] | None = None,
                fills=None) -> int:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    r = start_row
    for k, row in enumerate(rows):
        r = start_row + 1 + k
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=_cell_value(v))
            if decimals is not None and isinstance(cell.value, float):
                cell.number_format = _number_format(decimals[c - 1])
            if fills is not None and fills[k] is not None:
                cell.fill = fills[k]
    return r


def _autosize(ws, min_width: int = 8, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def _table_sheet(wb, title: str, res, key: str) -> None:
    df, cols, flag, muted = table_frame(res, key)
    if df is None:
        return
    ws = wb.create_sheet(title)
    labels = {**{c: column_label(c) for c in cols}, **dif_labels(res)}
    values = df[cols].to_numpy(dtype=object)
    fills = [FLAG_FILL if f else (MUTED_FILL if m else None) for f, m in zip(flag, muted)]
    _write_rows(ws, [labels[c] for c in cols], values.tolist(), decimals=[DECIMALS.get(c, 3) for c in cols],
                fills=fills)
    ws.freeze_panes = "B2"
    _autosize(ws)


def export_excel(res, interp, path: str | Path, source_name: str = "") -> Path:
    """Tulis workbook Excel berisi ringkasan dan semua tabel hasil analisis."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = S.SHEETS["summary"]
    ws.cell(row=1, column=1, value=S.REPORT_TITLE).font = Font(bold=True, size=14, color=_hex(theme.NAVY))
    ws.cell(row=2, column=1, value=generated_line())
    r = 4
    for label, value in metadata(res, source_name):
        ws.cell(row=r, column=1, value=label).font = BOLD
        ws.cell(row=r, column=2, value=value)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=S.SUMMARY_1MIN).font = Font(bold=True, size=12, color=_hex(theme.BRONZE))
    rows = [[lt.title, STATUS_LABELS[lt.status], lt.sentence] for lt in interp.lights]
    fills = [PatternFill("solid", fgColor=STATUS_FILLS[lt.status]) for lt in interp.lights]
    r = _write_rows(ws, S.LIGHTS_HEAD, rows, start_row=r + 1, fills=fills) + 2
    for sec in interp.sections:
        if not sec.technical:
            continue
        ws.cell(row=r, column=1, value=sec.title).font = Font(bold=True, size=12, color=_hex(theme.BRONZE))
        r = _write_rows(ws, S.TECH_HEAD, [[t.label.strip(), t.value, t.criterion] for t in sec.technical],
                        start_row=r + 1) + 2
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 90

    _table_sheet(wb, S.SHEETS["items"], res, "items")
    _table_sheet(wb, S.SHEETS["persons"], res, "persons")
    _table_sheet(wb, S.SHEETS["categories"], res, "categories")
    _table_sheet(wb, S.SHEETS["dif"], res, "dif")
    _table_sheet(wb, S.SHEETS["loadings"], res, "loadings")

    q3 = res.q3["matrix"]
    ws = wb.create_sheet(S.SHEETS["q3"])
    names = list(q3.columns)
    _write_rows(ws, ["Q3"] + names, [[n] + list(q3.loc[n].to_numpy()) for n in names],
                decimals=[0] + [3] * len(names))
    ws.freeze_panes = "B2"
    _table_sheet(wb, S.SHEETS["q3_pairs"], res, "q3_pairs")

    ws = wb.create_sheet(S.SHEETS["notes"])
    _write_rows(ws, ["Tingkat", "Pesan"], [[S.NOTE_LEVELS.get(i.level, i.level), i.message] for i in res.issues])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 140

    ws = wb.create_sheet(S.SHEETS["iterations"])
    _write_rows(ws, S.ITER_HEAD, [[rec.iteration, rec.max_change, rec.max_residual] for rec in res.jmle.log],
                decimals=[0, 6, 6])
    _autosize(ws)

    ws = wb.create_sheet(S.SHEETS["settings"])
    _write_rows(ws, ["Pengaturan", "Nilai"], [[k, str(v)] for k, v in res.settings.items()])
    _autosize(ws)
    wb.save(path)
    return path
