"""Ekstraksi teks dari DOCX.

Paragraf kosong diabaikan, sedangkan tabel dibaca baris demi baris karena pada
dokumen kebijakan dan catatan lapangan tabel kerap memuat data yang bermakna.
"""

from __future__ import annotations

from pathlib import Path

from src.ingestion._raw import resolve_raw
from src.ingestion.blocks import Block


def extract(path: Path | str) -> list[Block]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - bergantung lingkungan
        raise RuntimeError("Paket 'python-docx' belum terpasang.") from exc

    source = resolve_raw(path)
    dokumen = docx.Document(str(source))
    blocks: list[Block] = []

    for nomor, paragraf in enumerate(dokumen.paragraphs, start=1):
        teks = paragraf.text.strip()
        if teks:
            gaya = (paragraf.style.name if paragraf.style else "") or ""
            blocks.append(Block(
                text=teks,
                source_ref=f"par. {nomor}",
                meta={"style": gaya} if gaya else {},
            ))

    for nomor_tabel, tabel in enumerate(dokumen.tables, start=1):
        for nomor_baris, baris in enumerate(tabel.rows, start=1):
            sel = [c.text.strip() for c in baris.cells if c.text.strip()]
            if sel:
                blocks.append(Block(
                    text=" | ".join(sel),
                    source_ref=f"tabel {nomor_tabel} baris {nomor_baris}",
                ))
    return blocks
